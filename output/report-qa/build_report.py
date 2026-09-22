from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parents[1]
doc = Document()
s = doc.sections[0]
s.page_height, s.page_width = Inches(11.7), Inches(8.3)
s.top_margin = s.bottom_margin = Inches(.65)
s.left_margin = s.right_margin = Inches(.7)
styles = doc.styles
for name in ['Normal', 'Body Text', 'List Bullet']:
    styles[name].font.name = 'Calibri'
    styles[name].font.size = Pt(10.5)
    styles[name].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.08
for name in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']:
    styles[name].font.name = 'Calibri'
    styles[name].font.color.rgb = RGBColor.from_string('123D50')
styles['Title'].font.size = Pt(32)
styles['Heading 1'].font.size = Pt(22)
styles['Heading 2'].font.size = Pt(13)
styles['Heading 2'].paragraph_format.space_before = Pt(10)
for st in styles:
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
footer = s.footer.paragraphs[0]
footer.alignment = 2
r = footer.add_run('ORCA / SAMUDRA  |  ')
r.font.size = Pt(8)
f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), 'PAGE'); footer._p.append(f)
doc.core_properties.title = 'ORCA PS 26176: Prototype Gap Audit and Four-Role Implementation Plan'
doc.core_properties.subject = 'Codebase audit, stakeholder workflows, synthetic seeding and milestone plan'
doc.core_properties.author = 'SAMUDRA Project'

def p(text, style=None): doc.add_paragraph(text, style)
def h(text): doc.add_heading(text, 2)
def page(title):
    doc.add_page_break()
    doc.add_heading(title, 1)
def bullets(items):
    for t in items: p(t, 'List Bullet')
def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Light Shading Accent 1'
    if widths is None:
        widths = {2: [2.0, 4.9], 3: [1.2, 2.9, 2.8], 4: [2.7, .75, 2.7, .75], 5: [1.2, 1.425, 1.425, 1.425, 1.425]}[len(headers)]
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    for c, text in zip(t.rows[0].cells, headers): c.text = text
    rep = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(rep)
    for row in rows:
        cells = t.add_row().cells
        for c, text in zip(cells, row): c.text = str(text)
    for row in t.rows:
        pr = row._tr.get_or_add_trPr(); el = OxmlElement('w:cantSplit'); pr.append(el)
        for i,c in enumerate(row.cells):
            if widths: c.width = Inches(widths[i])
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(4)
                para.paragraph_format.space_before = Pt(3)
                for r in para.runs:
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor.from_string('202D35')
    return t

p('PS 26176  /  IMPLEMENTATION REVIEW')
doc.add_heading('ORCA\nPrototype gap audit', 0)
p('SAMUDRA • Marine EcOsystem Reasoning with Collaborative Agents', 'Subtitle')
p('Detailed implementation report and fresh four-role delivery plan', 'Subtitle')
p('18 September 2026')
h('Executive assessment')
p('SAMUDRA contains a substantial conversational and visualization foundation, but the current wiring does not yet establish a fully functional prototype across the supplied problem statement. Three stakeholder dashboards and a deterministic synthetic dataset already exist. The highest-value remaining work is to connect those surfaces to consistent, time-aware domain results, complete proactive alert workflows, and support the missing stakeholder journeys.')
bullets([
    'Preserve: Fisher, Authority and Researcher pages; chat and voice components; evidence and trace views; deterministic risk, PFZ and geofence modules; data normalizers; scenario framework; seed tooling.',
    'Fix first: provider-mode registration of mock PFZ/geofence/route tools, fixed route metrics, location/time fallback semantics, and misleading frontend outage fallbacks.',
    'Complete: Disaster Management and Maritime Operator dashboards, forecast-aware what-if evaluation, evidence-backed EO/productivity analytics, and durable alert delivery with replay-based geofence approach events.',
    'Deliver through four existing workstreams: M1 Experience, M2 Platform/Data, M3 Agents, M4 Marine/GIS/Risk. Five shared milestone gates span an indicative 20-working-day plan.'
])
h('Audit baseline')
p('Branch: main. Commit: bf8e15a12286080a8a9c76f9eefd99c3ea0e3e9a. The working tree was clean at the start. This report proposes implementation work; it does not change application code or supersede accepted project decisions.')
p('Scope: the supplied PS text, canonical repository documents, application source, tests, fixtures, migrations and deployment configuration. Status labels distinguish code present from end-to-end verification. This is a prototype assessment, not certification for operational navigation.')

page('1. Findings and verification boundary')
h('What was checked')
p('Reviewed the repository entry rules, ownership, API/data/safety contracts, master context, progress, decisions and existing implementation plan. Traced UI → client → API → AgentRunService → graph → registry → provider/domain paths, plus the synthetic generator, seeder and researcher/authority data paths. Source anchors appear in Section 16.')
table(['Check performed', 'Current result'], [
('Repository state', 'main at bf8e15a; no pre-existing working-tree changes.'),
('Python syntax', '154 backend and test Python files parsed successfully with ast.parse. This is a syntax check, not execution coverage.'),
('Synthetic generator', 'Executed via runpy without importing the application package. Two consecutive generations were equal; 1,024 records across 15 entity groups.'),
('Targeted backend tests', 'python -m pytest tests/domain/test_route_balanced.py tests/domain/test_marine_source_of_truth.py tests/test_synthetic_seed.py -q stopped during conftest import: missing shapely. No test assertions ran.'),
('Frontend tests', 'npm.cmd test -- --run failed because Vitest was unavailable; frontend/node_modules is absent. No frontend pass count is established.'),
('Runtime verification still required', 'Browser journeys, TypeScript/build, PostGIS persistence, live source access, voice providers and full regression must be rerun in a provisioned environment.')
])
h('Classification used throughout')
p('Present: source exists and the relevant path was inspected. Partial: useful behavior exists but a PS workflow or integration remains incomplete. Missing: no corresponding implementation was located in the inspected application paths. Unverified: execution/access must still be demonstrated. These labels are not percentages of completion.')
h('Important conflicts to resolve before coding')
bullets([
    'PROGRESS describes evaluated route alternatives and guaranteed water-only corridors; MockRouteExposureEngine still assigns fixed exposure values and uses a westward longitude clamp rather than a coastline intersection proof.',
    'SYNTHETIC_DATA documents 644 records and two harbors; current generation returns 1,024 records and five harbors. It recommends SYNTHETIC mode, while Settings.validated_data_mode accepts only LIVE/HYBRID/SNAPSHOT.',
    'SAFETY requires unconditional active hazard hard-stops; risk_engine returns UNKNOWN for degraded inputs before evaluating the cyclone flag. Missing/malformed validity also needs explicit conservative handling.',
    'README/master context retain old counts and a non-existent ORCA_IMPLEMENTATION_PLAN.md link. DECISIONS repeats D021; PROGRESS repeats milestone labels and mixes historical counts.'
])
p('These conflicts are reported, not silently resolved. Milestone G0 must record the approved contract and ownership decisions before affected implementation proceeds.')

page('2. Problem-statement coverage')
table(['PS requirement', 'Assessment and remaining implementation'], [
('Natural-language intent and task planning', 'Partial. LangGraph, intent classification, optional LLM planning, tool ordering and traces exist. Add explicit compound-query decomposition, EO/productivity tools, capability discovery by coverage, and mandatory safety dependencies.'),
('Same-language, multi-turn dialogue', 'Present foundation. EN/HI/MR UI and memory exist; Tamil is partial across layers. Verify language changes, role/sector isolation, persistent context and localized failure responses.'),
('Nearest PFZ today', 'Partial. Real ranking code exists, but provider graph registration can retain MockPFZRankingEngine with fixed candidates and empty raw_features. Wire real raw advisories → validation → ranking → safety exclusion.'),
('Safe tomorrow morning', 'Partial. Deterministic thresholds exist. Departure time is not consistently propagated into observation selection; snapshot marine selection favors the fixed reference record.'),
('Tides, weather and sea state', 'Partial. Researcher tide/sea-state charts exist. Add typed tide results and requested-time/station retrieval to conversational tools and evidence.'),
('Lightning and cyclone alerts', 'Partial. Hazard fixtures, connectors and maps exist. Add explicit lightning coverage semantics, conservative outage handling and proactive delivery.'),
('Favorable chlorophyll and SST regions', 'Partial. EO normalizer, 350 seeded grid cells and charts/maps exist. Add registered analytical tools to correlate valid cells by region/time and explain candidate selection.'),
('Safest route / route optimization', 'Partial. Candidate UI and route interfaces exist; runtime route metrics remain simulated constants. Evaluate traversed edges against time-indexed conditions, restrictions and vessel limits.'),
('Why fish productivity declined', 'Missing dedicated evidence workflow. Generic analytical explanation is not a productivity model. Add historical environmental/catch-effort evidence and qualified hypotheses, not causal claims from chlorophyll alone.'),
('Avoid hazardous / restricted zones', 'Partial. Real GIS checks exist, but graph uses a mock geofence adapter. Add safe PFZ exclusion, approach notifications, temporal restrictions and actual-route checking.'),
('Autonomous source discovery/integration', 'Partial. Fixed adapters and normalizers exist. Add an allowlisted source/product catalog with variable, spatial, temporal, quality and access metadata; no unrestricted web agent required.'),
('Maps, charts, evidence and reasoning', 'Strong UI foundation. Close mismatch between displayed and evaluated data, preserve provenance on derived results, and attach alert/mission events to execution IDs.')
])
h('Prototype scope versus later expansion')
p('A functional demo can use clearly labeled synthetic observations and versioned public-data snapshots, provided the actual algorithms compute the results. Genuine public EO ingestion should be demonstrated with at least one licensed/accessible sample and its acquisition metadata. Continuous government feeds, AIS hardware, SARAT, satellite messaging and national coverage are later integration work, not prerequisites for an honest controlled prototype.')

page('3. Priority defects and integration gaps')
table(['Priority / evidence', 'Required correction and acceptance condition'], [
('P0: mock domain wiring [E2–E4]', 'Provider registration installs marine/weather/hazard/SVAS, then fills missing capabilities with contract mocks. Register actual PFZ and GIS engines; isolate mocks to explicit tests/scenarios. A provider-mode test must assert implementation identity and changed output when input data changes.'),
('P0: route constants [E4]', 'API directly constructs MockRouteExposureEngine. Replace fixed wave maxima, scores and always-inshore recommendation with geometry/time-dependent evaluation. Preserve the current route interface where possible; return no feasible route when every option is blocked.'),
('P0: location substitution [E5]', 'SnapshotConnector falls back to Ratnagiri records when another harbor lacks observations, then assigns the requested harbor label. Return unavailable coverage or a disclosed spatial interpolation with distance/uncertainty. Never relabel another station as local.'),
('P0: time and freshness [E5–E7]', 'Pass absolute departure/window and simulation reference time end to end. Snapshot selection must use the requested time, not hour_offset == 0. Missing or malformed timestamps cannot establish freshness.'),
('P0: outage data fabrication [E8]', 'fetchOrMock returns fixed marine, EO, PFZ, hazard and healthy/connected records after HTTP errors. Fisher also consumes PFZ/hazard fetches. Replace operational fallbacks with explicit unavailable state or an explicitly selected versioned offline dataset.'),
('P0: hazard semantics [E6]', 'Do not encode absent bulletins as NORMAL with newly minted validity. Preserve a valid known hard-stop when other feeds fail; otherwise disclose unknown coverage. Validate these outcomes after resolving the safety-contract conflict.'),
('P1: real research tools [E9]', 'Tide, EO and productivity analysis need normalized result contracts, deterministic computations, registered capabilities and evidence-linked response templates.'),
('P1: durable workflows [E10–E12]', 'Demo endpoints are predominantly read-only. Add saved missions, subscriptions, incident/alert lifecycle and tenant/role boundaries; persistent SQL models alone do not supply these workflows.')
])
h('Do not solve these through cosmetic changes')
p('More 3D styling, additional canned prompts, or renamed mock classes will not close the PS gaps. Retain the modular monolith and focus on end-to-end data lineage and behavior. A named “specialist agent” is useful only when it selects and executes a real capability with inspectable inputs, outputs and evidence.')

page('4. Stakeholder dashboards: retain and complete')
p('The PS names five stakeholder groups. Separate dashboards are a requested product requirement; they do not require five separate applications. Use five role routes with a shared UI kit, map engine, evidence drawer and typed API clients. The existing three portal modes should be extended, not replaced.')
table(['Dashboard', 'Current surface', 'Completion work / demo acceptance'], [
('Fisher / skipper', 'FisherPage; decision surface, local map, chat, call UI, mission context and what-if.', 'Persist vessel/harbor preferences and a trip plan; show tide + requested forecast time; rank safe PFZs; deliver approach/hazard alerts; low-bandwidth stale state. Demo: change vessel or departure and see a recomputed, cited decision.'),
('Coastal authority / harbor master', 'AuthorityPage; sectors, fleet replay, hazard containment, operational alerts, routes and audit.', 'Add advisory lifecycle, acknowledgements, fleet filters and incident handoff. Explain playback time versus latest evaluated position. Demo: select alert → vessel + hazard → evidence/run → acknowledged record.'),
('Marine researcher', 'ResearcherPage; ocean explorer, source monitor, scenario lab and query workbench; EO/tide charts.', 'Add region/date selection tied to tools, reproducible EO correlations, productivity hypotheses, exports and actual source health. Propagate app locale. Demo: regenerate an analytical result from dataset IDs and selected period.'),
('Disaster management', 'Seeded stakeholder identity; no dedicated App portal mode.', 'New DisasterPage: incident queue, affected vessels/sectors, warning timeline, response assignment and acknowledgement status. Reuse Authority maps; do not imply real rescue dispatch. Demo: escalating synthetic cyclone creates an incident and updates affected assets.'),
('Maritime operator', 'Seeded identity plus fleet/trip data; no dedicated App portal mode.', 'New OperatorPage: own-fleet trips, schedule, departure windows, route comparison, mission status and exceptions. Demo: save a voyage, reject a hazardous route, delay departure and compare updated exposure.')
])
h('Common behavior for every dashboard')
bullets([
    'Role-appropriate landing page, explicit data mode/reference clock, location and time window, source quality, and consistent loading/empty/unavailable states.',
    'URL-based navigation/deep links and per-role conversation scope. Current App shares a useChat instance across Fisher and Authority; do not let hidden prior-role context drive a new role.',
    'Server-side role/organization checks for writes and data scope. A portal card or client-side “logout” is not authentication.',
    'Keyboard access, responsive layouts, readable safety language and reduced-motion mode. All roles may share facts; none may recompute authoritative safety in the browser.'
])

page('5. Synthetic data: current inventory and gaps')
p('Fake-data seeding is already a significant feature. Extend the existing generator → normalizer → validator → repository path. Synthetic inputs are acceptable; hardcoded final recommendations and fabricated outage telemetry are not.')
table(['Entity group', 'Count', 'Entity group', 'Count'], [
('Stakeholders',5,'Harbors',5),('Fishers',8,'Vessels',14),('Trips',18,'Marine observations',96),
('EO grid cells',350,'PFZ candidates',12),('Geofences',5,'Route nodes',24),('Route edges',32,'Hazards',10),
('Notifications',20,'Replay positions',420),('Sectors',5,'Total generated records',1024)
])
p('Verified by direct generator execution, not a database seed. Reference time is 2026-09-12 06:00 UTC. Marine observations cover two harbors for 48 hours; five fleet sectors do not imply five complete environmental coverage areas. EO data covers a 5×5 grid over 14 days.')
h('Required extensions')
bullets([
    'M4: add named reproducible scenarios for calm, rough sea, cyclone escalation, lightning, restricted-zone approach/crossing, expired PFZ, cloud gaps, missing tide, provider conflict and complete outage. Store scenario ID, version and reference clock.',
    'M4: add time-varying edge conditions, valid/invalid route candidates and independently checked water geometry. Seed inputs so route rankings can change with departure time, craft and hazard activation.',
    'M4: add at least 30 daily catch-and-effort records for two pilot regions, with species/gear/effort units, nulls and confounding events, solely to demonstrate qualified productivity analysis.',
    'M2: persist role-to-vessel/sector scope, saved missions, alert state/acknowledgements and incident assignments. Seed synthetic identifiers only; do not import real fisher contact details.',
    'M2: reconcile sectors with persistence. The generator includes sectors, but seeder.entity_model_map has no sector model. Ensure generated, DB-backed and file-backed API views agree.',
    'M2/M4 interface: repair custom-namespace mapping for nested links such as affected_sector_ids and validate all references. Keep the seeder CLI; do not build a second seeding framework.'
])
h('Seed acceptance gate')
p('Run dry-run validation; seed twice with identical counts and no duplicates; reset only a disposable target namespace; verify unrelated namespaces survive; compare generated fixtures, DB records and API responses by IDs and checksums. Every screen must preserve mode, observation time, validity, units, QC and reference clock. Test with the wall clock advanced: frozen demo mode remains reproducible, while real-time mode rejects expired records.')

page('6. Target prototype behavior and contracts')
h('One computation path')
p('Source adapters or explicit demo datasets → canonical observation/evidence store → M4 analytical/risk/GIS functions → M3 bounded orchestration → M2 response and persistence → M1 role dashboard. The alert worker and chat must call the same domain functions. Research charts and chat must use the same dataset IDs and time windows.')
table(['Proposed contract', 'Minimum fields and behavior', 'Single editor'], [
('Query / mission context', 'mission_id, role scope, vessel_id, origin/destination geometry, departure_at, duration, timezone, dataset_version and optional simulation clock. Invalid or uncovered context returns a typed error.', 'M2'),
('Source coverage', 'product_id, variables/units, spatial footprint, time extent, update cadence, access status, mode, last successful fetch, QC and provenance. Distinguish source access from data freshness.', 'M2'),
('Route result', 'candidate geometry, computed distance/ETA, sampled time range, exposure basis, hard-stop exclusions, feasibility, evidence IDs and no-feasible-route outcome.', 'M2; M4 specifies semantics'),
('Research result', 'region/window, valid/missing counts, statistics, units, uncertainty, method version and evidence IDs. Productivity answers distinguish observation, association and hypothesis.', 'M2; M4 specifies semantics'),
('Alert / incident', 'stable event ID, vessel/sector/hazard IDs, observed/effective/expiry times, trigger type, severity, lifecycle, run/evidence linkage, delivery state and acknowledgement actor/time.', 'M2'),
('Agent plan / trace', 'subtasks, selected capabilities, dependencies, execution status, timing and evidence references. Retain concise rationale; do not expose hidden model reasoning.', 'M3 within agents; M2 public schema')
])
h('Safety and time rules to ratify at G0')
p('Valid known prohibitions and active severe warnings remain hard-stops even if a non-decisive observation is missing. Unknown essential coverage prevents GO. Evaluate restrictions for the voyage time, not only machine time. Derive confidence from real provenance and quality; synthetic provenance must not become official-source confidence. National EEZ/territorial outlines are informational layers, not automatically forbidden waters.')
h('Minimal proactive alert loop')
p('Consume replay/position or provider updates → validate freshness/sequence → evaluate approach/containment against effective hazards and boundaries → deduplicate event → persist → publish/poll → display/audio → acknowledge. A prototype may use local replay and polling; SMS, VHF or push-provider integration is optional. Acceptance requires an alert without a new chat question and without flooding duplicates at each replay tick.')

page('7. Four roles and non-overlapping ownership')
p('Keep the existing M1–M4 team split. Five stakeholder dashboards are product surfaces, not five engineering roles. Each file and deliverable below has one writer. Other members provide specifications or review; they do not edit another workstream’s files.')
table(['Role', 'Exclusive write scope', 'Handoff boundaries'], [
('M1 · Experience', 'frontend/**, frontend tests and frontend/src/types/contracts.ts.', 'Consumes published API/domain results. Owns all five dashboard UXs, accessibility and browser tests. No backend risk or connector code.'),
('M2 · Platform / Data', 'Existing API, contracts, connectors, repositories, core, migrations/deployment. Proposed explicit extension: db/**, services/**, main.py, scripts/seed_demo.py and root CI.', 'Publishes schemas and persistence/services. No agent prompts or domain formulas. Owns canonical API/data documentation edits and final release assembly.'),
('M3 · Agents / AI', 'backend/app/agents/**, prompts/**, agent evaluation tests. Proposed explicit extension: scenarios/runner.py and scenario execution models.', 'Registers and invokes M4 functions through approved interfaces. Owns conversation context, planning, localization and evidence composition.'),
('M4 · Marine / GIS / Risk', 'backend/app/domain/**, tools/**, data/fixtures/** and domain/scenario fixture tests. Proposed extension: data/reference/** and scenario registry/fixture definitions.', 'Provides deterministic functions, scientific semantics and data packages; no HTTP connector or UI edits. Owns safety-document edits.')
])
h('Assignments that need an integration decision')
p('OWNERSHIP.md leaves db/, services/, scripts/, main.py and scenario execution ownership implicit. The extensions above are proposals, not assumed permission. M2 records the approved assignments and shared interface decisions in existing DECISIONS.md and OWNERSHIP.md at G0. M3 remains the editor of existing agents/integrations/dev2.py and dev4.py until an explicitly agreed interface move; M4 supplies the contract semantics.')
h('Shared-file protocol')
bullets([
    'M2 is sole editor of API_CONTRACTS, CANONICAL_DATA_CONTRACTS, IMPLEMENTATION_PLAN, PROGRESS, OWNERSHIP and DECISIONS during integration; M4 edits SAFETY. M1 alone edits TypeScript consumer types.',
    'Each role writes its own tests. M2 owns API/DB/connector integration tests; M3 owns graph and scenario-runtime integration tests; M4 owns numerical/GIS/data validation tests; M1 owns UI and browser tests.',
    'Publish additive contracts and golden payloads before consumers start. Consumer work uses explicit fixtures until producers pass contract tests. No shared-file cherry-picking across simultaneous branches.',
    'Use the documented team/* → develop → main flow. Integrate daily; freeze schemas at each gate. “No overlap” means no competing implementation or file ownership, not zero collaboration.'
])

page('8. M1 milestones: experience and dashboards')
p('M1 owns the entire frontend so shared shell, maps and styles stay coherent. Reuse existing dashboards and primitives. The two new pages are thin consumers of M2 services, not duplicated operational backends.')
table(['ID / gate', 'Task and dependency', 'Acceptance evidence'], [
('U1 · G0', 'Audit five stakeholder journeys; lock role routes, shared layout and state diagrams. Input: agreed scope/contracts.', 'Five role-to-action maps and explicit empty/error/loading states; current three dashboards retained.'),
('U2 · G1', 'Remove operational fetchOrMock fallbacks; consume canonical mode/health/coverage. Dependency: P2 response envelope.', 'Backend-offline browser test shows unavailable, never healthy/connected or invented PFZ/hazard data.'),
('U3 · G1', 'Add URL routing, scoped sessions and role navigation; isolate mission context when switching roles/sectors. Dependency: P1 identity contract.', 'Direct links and refresh work; switching role/sector cannot display a previous role’s active recommendation.'),
('U4 · G2', 'Complete Fisher trip/time/route/PFZ workflows and real decision-diff UI. Dependencies: P3, A3 and D3.', 'Departure/craft/destination changes update results; no feasible route is explicit; displayed evidence matches map geometry.'),
('U5 · G2', 'Complete Researcher region/date queries, analytical results and reproducible export view. Dependencies: P4, A4 and D4.', 'Chart/table/chat use matching sample window and dataset IDs; QC gaps survive export; no unsupported causal wording.'),
('U6 · G3', 'Extend Authority alert audit/acknowledgement and build Disaster dashboard using shared fleet/hazard components. Dependency: P5 and D5.', 'Replay creates an alert; operator inspects evidence, acknowledges it and sees the updated incident state after refresh.'),
('U7 · G3', 'Build Maritime Operator dashboard: own fleet, saved trips, schedule and route comparison. Dependency: P3/P5.', 'Operator sees only allowed fleet; saves a mission and re-evaluates after a forecast change.'),
('U8 · G4', 'EN/HI/MR localization, voice fallback, keyboard/mobile/reduced-motion QA and five-role browser suite.', 'All role journeys pass; microphone denial works; status semantics survive language changes; npm test/typecheck/build succeed.')
])
h('Out of scope for M1')
p('No client-side safety scores, source freshness invention, simulated operational telemetry after failed requests, or hardcoded “best route.” Turf geometry may support presentation/inspection, but the backend remains authoritative for decisions.')
h('Estimated capacity')
p('Approximately 18–20 focused person-days within the proposed four-week window. Keep 3D enhancements frozen until U2–U7 pass. If capacity slips, simplify disaster/operator layouts using shared components rather than cutting the data integrity work.')

page('9. M2 milestones: platform and data')
table(['ID / gate', 'Task and dependency', 'Acceptance evidence'], [
('P1 · G0', 'Reconcile mode/ownership/contract conflicts; publish identity, query context, source coverage, route, research and alert schemas. Assign shared-file editors.', 'Approved decisions; one schema per concept; backward-compatibility and field ownership documented.'),
('P2 · G1', 'Correct SnapshotConnector coverage/time semantics, unavailable hazards/SVAS provenance and dynamic provider health. Supply typed source catalog.', 'No Ratnagiri relabeling; no invented timestamps; each record retains source/mode/QC; outage responses distinguish unavailable from empty.'),
('P3 · G1–G2', 'Complete PostGIS migration coverage, transactional seed/reset and saved mission APIs; add minimal demo auth and server-side scope checks.', 'Clean DB upgrade, seed twice, restart/read and downgrade rehearsal; unauthorized cross-role/organization reads/writes rejected.'),
('P4 · G2', 'Expose time-windowed observations, EO/tide/productivity inputs and analytical endpoints. Import one actual public EO sample when access permits.', 'Contract fixtures match M4 functions; public sample has acquisition metadata/checksum/license; synthetic fallback is separately labeled.'),
('P5 · G3', 'Implement alert worker/event persistence, dedupe, subscriptions, delivery polling/stream, acknowledgements and incident lifecycle. Dependency: D5 trigger evaluator.', 'No chat action needed; idempotent replay; state survives restart; role-scoped acknowledgement audit is preserved.'),
('P6 · G3', 'Add export endpoints and end-to-end run/evidence linkage for mission/alert/research outputs; typed errors and bounded pagination.', 'Export can be traced to a saved run; incomplete data and access failures return stable errors.'),
('P7 · G4', 'Provision reproducible backend/frontend dependencies, CI, complete Compose deployment, readiness and logs. Declare sarvamai dependency or optional voice extra.', 'Fresh checkout installs and starts via documented commands; checks run in CI; DB migrations, static frontend and API form one deployable demo.'),
('P8 · G4', 'Integrate release, update existing canonical progress/plan, collect all owners’ test evidence and run provider-outage/restart rehearsal.', 'No P0 defects open; no secret values in logs/artifacts; bounded request sizes/timeouts and release checklist pass.')
])
h('Implementation cautions')
p('The repository already has SQLAlchemy models, offline persistence, a Dockerfile and an initial Alembic migration. Extend them rather than declaring the database missing. Startup create_all currently masks migration completeness. Compose activates only the DB; backend/frontend services are commented out. Use migration tests to establish parity with all demo/mission/alert tables.')
p('Provider URLs in configuration include placeholders. Connector code is not proof of access. Maintain source-specific status and allow a release with disclosed historical/demo coverage; do not claim LIVE solely because a model or key is configured.')

page('10. M3 milestones: agents and reasoning')
table(['ID / gate', 'Task and dependency', 'Acceptance evidence'], [
('A1 · G0', 'Map actual tool registry to PS questions and ratify mandatory safety capabilities, compound-query output and context fields with P1.', 'Capability-to-question matrix; safety prerequisites cannot be dropped by an LLM plan.'),
('A2 · G1', 'Replace provider-mode mock domain handlers with explicit bindings to D2/D3; separate scenario/test registries from application runtime.', 'Provider-mode PFZ/geofence tools call real engines; concurrent scenario execution cannot change a normal chat registry or memory store.'),
('A3 · G2', 'Propagate location, absolute departure/window, objective and simulation clock through intent, memory and tool invocation. Implement structured what-if comparison.', 'Delay changes sampled forecasts; craft change reevaluates limits; original mission/context preserved; unsupported time horizon returns unavailable.'),
('A4 · G2', 'Register tide, EO suitability and productivity-analysis tools; add multi-source compound plans using P2 catalog and D4 functions.', 'Queries combine SST/chlorophyll/weather/route evidence; productivity response identifies missing catch/effort evidence and states qualified hypotheses.'),
('A5 · G2', 'Unify actual planner capability names with supervisor contracts; bounded retry/replan on missing coverage, no unbounded autonomous fetch loops.', 'Trace records selected source/capability and failure reason; tool count/latency budget enforced; missing evidence cannot be hidden in fluent text.'),
('A6 · G3', 'Ground explanations of changed recommendations, alert triggers and mission alternatives; link tool evidence to returned maps/charts.', 'Every quantitative conclusion has evidence IDs and units; generated narrative cannot contradict deterministic status or route exclusion.'),
('A7 · G3', 'Finish same-language behavior for EN/HI/MR; evaluate Tamil as explicit optional expansion; scope persistent memory by actor/role/mission.', 'Mid-conversation language switch works; translated hazard terms preserve severity; role and mission context do not leak.'),
('A8 · G4', 'Expand PS acceptance/evaluation corpus, injection/unsupported-claim tests, provider/LLM outage behavior and multi-user isolation.', 'All eight supplied example query families execute against declared fixtures; deterministic fallback remains truthful; CI evaluation results are published.')
])
h('Agent design boundary')
p('Keep one bounded graph with specialized capabilities. The current sequential LangGraph pipeline is useful and does not need a microservice rewrite or one LLM per specialist. Add planning flexibility only where compound questions and data gaps require it. M4 owns arithmetic, scientific aggregation, geometry and safety decisions.')
h('Critical integration test')
p('Use the same mission twice, change only the underlying forecast or hazard fixture, and verify that tool outputs, recommendation, explanation and map all change consistently. This is stronger evidence of reasoning than a canned scenario returning an expected status through a test double.')

page('11. M4 milestones: marine, GIS and risk')
table(['ID / gate', 'Task and dependency', 'Acceptance evidence'], [
('D1 · G0', 'Resolve safety precedence, clock, units, confidence, coverage and restriction semantics with P1. Define bounded pilot geography.', 'Approved rules for hard-stop plus missing-data cases, timestamp failures and known/unknown source coverage.'),
('D2 · G1', 'Harden existing risk/PFZ/GIS engines: valid-time filtering, QC, craft radius, restricted/hazard exclusions, missing origin and accurate proximity.', 'No fabricated origin; known prohibition cannot become GO/UNKNOWN solely due to unrelated missing data; expired/future restrictions handled; deterministic ranking verified.'),
('D3 · G2', 'Implement a real route evaluator over the existing graph: geodesic segment distances, edge traversal times, wave/wind/current samples and hard constraints.', 'Hazard/forecast/departure perturbations change feasible routes and ranking; all blocked returns no feasible route; coast/land checks replace longitude-only assumptions.'),
('D4 · G2', 'Implement tide lookup, EO QC-aware aggregation/suitability and catch-effort trend comparison with uncertainty and evidence lineage.', 'Null is not zero; cloud gaps excluded explicitly; units/datum preserved; correlation is not reported as proven cause.'),
('D5 · G3', 'Implement pure approach/entry/exit hazard and boundary evaluators with buffer/hysteresis and event-time semantics.', 'Replay triggers approach before entry, dedupes jitter, expires warnings, rejects stale/out-of-order points; no worker/persistence logic in domain.'),
('D6 · G3', 'Extend canonical synthetic scenarios, route/forecast inputs and productivity datasets; validate geometry and nested relationships.', 'Golden manifests/checksums stable; bad geometry and broken relations fail; data covers each PS workflow and each target role.'),
('D7 · G3', 'Create independent numerical oracles and adversarial boundary cases for routing, units, tide and risk; review source semantics from P4.', 'Threshold edges, intersecting restrictions, coastal bends, missing samples and all-routes-blocked cases have meaningful assertions.'),
('D8 · G4', 'Run domain/scenario regression and review prototype scientific/safety claims; update SAFETY after approved decisions.', 'Measured limitations disclosed, reproducible results supplied to release, no unsupported water-only or operational-safety claims.')
])
h('Refactoring boundary')
p('Route logic currently lives inside agents/integrations/mocks.py, which belongs to M3. M4 creates the real domain implementation behind the accepted interface; M3 alone changes the old binding. Do not move or rename mock code into domain and call it complete.')
h('Geographic scope')
p('Validate Ratnagiri–Malvan/Konkan deeply first. Other fleet sectors may remain disclosed replay-only demonstrations until local observation and coastline coverage exist. A west-of-harbor rule cannot support arbitrary Indian coasts or islands.')

page('12. Integrated milestone calendar')
p('Indicative estimate: 20 working days, four focused contributors, approximately 72–80 person-days including integration. Day 1 begins after G0 decisions and environment setup are available. This is a planning estimate, not a guaranteed deadline; real-provider access can run in a separate external dependency lane.')
table(['Gate / window', 'M1', 'M2', 'M3', 'M4'], [
('G0 · Days 1–2\nBaseline locked', 'U1: journeys / role shell contract', 'P1: contracts, ownership and baseline', 'A1: capability and safety plan', 'D1: safety/time semantics'),
('G1 · Days 3–6\nTruthful foundation', 'U2–U3: honest states / scoped roles', 'P2 + P3 foundation: canonical reads / seed / scope', 'A2: real domain bindings / isolation', 'D2: risk / PFZ / GIS corrections'),
('G2 · Days 7–11\nPS reasoning', 'U4–U5: Fisher / Researcher workflows', 'P3 completion + P4: missions / research data', 'A3–A5: temporal / research / compound plans', 'D3–D4: routes / tide / EO / productivity'),
('G3 · Days 12–16\nFive-role prototype', 'U6–U7: Authority / Disaster / Operator', 'P5–P6: alerts / incidents / exports', 'A6–A7: grounded diffs / locale / memory', 'D5–D7: triggers / datasets / oracles'),
('G4 · Days 17–20\nRelease candidate', 'U8: browser / mobile / locale QA', 'P7–P8: deployment / CI / release', 'A8: PS evaluations / resilience', 'D8: regression / safety review')
])
h('Gate exit criteria')
bullets([
    'G0: conflicts recorded, ownership approved, reproducible dependency setup, current test baseline captured, additive schemas and golden payloads published.',
    'G1: operational outage cannot fabricate success; provider mode binds actual domain engines; canonical location/time/mode survive across API and UI.',
    'G2: all eight PS query families have an executable path, including EO/productivity and tide; temporal/route perturbation tests prove recomputation.',
    'G3: five dashboards perform distinct end-to-end workflows; replay causes deduplicated alerts; mission/acknowledgement/incident state persists.',
    'G4: full checks pass, fresh-machine demo works, no P0 defects remain, provenance and safety claims match behavior, browser and restart/outage rehearsal recorded.'
])
h('Critical dependencies and integration order')
p('P1 → D2/P2 → A2 → U2/U4 is the foundation chain. P1 → D3/P4 → A3/A4 → U4/U5 closes PS reasoning. D5 → P5 → U6/U7 closes proactive operations. Publish contracts first so M1 can work with explicit golden payloads; integrate producers before removing those development-only fixtures.')
p('If time is constrained, reduce decorative UI, national coverage and external delivery channels. Do not cut truthful degraded states, real domain wiring, hard-stop tests or reproducible seed behavior. Do not label a three-page mock-only experience a complete five-stakeholder prototype.')

page('13. Focused refactoring plan')
table(['Change / owner', 'Evidence and smallest useful scope', 'Regression guard'], [
('R1 · M3: runtime bindings', 'graph.py re-registers global tools; provider mode fills with contract mocks; scenarios also use global registry/memory. Introduce explicit runtime/scenario dependency containers or factories without replacing LangGraph.', 'Concurrent scenario/chat requests retain their intended providers, state and evidence.'),
('R2 · M4 + M3 handoff: real routes', 'Move new calculations into domain, with M3 switching adapter bindings. Fixed metrics and westward clamps are not route optimization.', 'Input perturbation changes output; independently recomputed distance/feasibility agrees.'),
('R3 · M2: canonical data access', 'SnapshotConnector imports API _model_to_dict and mixes DB/file paths; demo APIs and DataService also select data. Extract shared serialization/query service at the platform boundary.', 'Same location/time/version produces same canonical records for chart, chat and API.'),
('R4 · M1: fallbacks and types', 'researcher-client contains operational MOCK_* records and broad any mappings. Use typed result/error envelopes and one explicit offline source selection.', 'Failed fetches cannot report healthy DB or fabricated advisory values; null/QC fields survive mapping.'),
('R5 · M2: API module boundaries', 'routes.py combines chat, voice, scenarios and many demo endpoints in one large file. Split by existing endpoint responsibility only while adding related workflow APIs.', 'OpenAPI paths/status codes and response schemas unchanged unless explicitly versioned.'),
('R6 · M3: planner contracts', 'supervisor.py standard names differ from graph capability names. Establish one capability catalog and dependency policy; remove redundant plan tables after consumers migrate.', 'Every declared capability resolves and mandatory safety dependencies remain enforced.'),
('R7 · M1: dashboard reuse', 'Two new roles will otherwise duplicate Authority maps and fleet views. Extract reusable components only where used by multiple roles; keep persona-specific actions separate.', 'Selection, alert focus, localization and degraded states work across roles.'),
('R8 · M2: canonical documentation', 'Old counts/paths, repeated decision IDs and incompatible status claims obscure the baseline. Consolidate existing docs after verified milestones.', 'One current plan/status board; no new competing roadmap markdown files.')
])
h('Refactors to avoid')
p('Do not replace React, FastAPI, LangGraph or the modular monolith. Do not introduce microservices, a second seeder, an additional map engine, a second risk system or unnecessary repositories. Do not rewrite all CSS or split files merely for line count. Tie each change to the concrete defect or duplication above.')
h('Change control')
p('Contract or ownership changes precede implementation. Keep each PR a coherent slice with purpose, files, contract impact, tests, risks and dependencies. M2 coordinates integration documentation; each role edits only its assigned files.')

page('14. Acceptance and test matrix')
table(['Scenario / owner', 'Required visible and technical outcome'], [
('Nearest safe PFZ · M4 + M3', 'Change origin and verify independent distance/bearing calculation; reject expired/out-of-range/restricted candidates; returned map IDs match evidence.'),
('Tomorrow safety / what-if · M4 + M3', 'Same vessel at two departure times samples different forecast intervals; no data beyond forecast horizon returns UNKNOWN; explanation identifies changed factors.'),
('Tide and sea state · M4 + M1', 'Requested location/time selects correct station; tide datum and units displayed; missing tide stays missing; chat and chart agree.'),
('Cyclone/lightning outage · M4 + M2', 'Known effective severe warning remains hard-stop despite a separate feed outage; unknown lightning coverage is not “no lightning”; stale warnings are not relabeled fresh.'),
('EO suitability · M4 + M3', 'Region/date query aggregates only appropriate valid cells; threshold/method and uncertainty recorded; evidence supports each reported number.'),
('Productivity decline · M4 + M3', 'Compare catch per effort and environmental series; disclose missing effort/seasonal context; separate correlations from causal hypotheses.'),
('Safest route · M4', 'Change one edge forecast/restriction and ranking responds; compute distances from route geometry; reject land/prohibited crossings; all blocked returns no feasible route.'),
('Proactive boundary alert · M2 + M4', 'Replay produces approach/entry event without chat; duplicates/jitter suppressed; acknowledgement and incident state persist and are role-scoped.'),
('Five dashboards · M1', 'Browser tests exercise real API flows for Fisher, Authority, Researcher, Disaster and Operator; deep links, sector switching, offline states and export work.'),
('Persistence / seed · M2', 'Alembic fresh DB + seed + restart + replay retain expected records; second seed is idempotent; namespace reset cannot erase another demo.'),
('Language / voice · M1 + M3', 'EN/HI/MR safety status preserved across text/audio and language switches; provider failure and denied microphone retain usable text flow.'),
('Security / concurrency · M2 + M3', 'Role/tenant isolation, prompt injection, unsupported numbers, scenario/chat overlap, cancellation and provider failure cannot contaminate another run.')
])
h('Checks before release')
p('Provision dependencies, then run the smallest relevant domain/contract test per change. Run affected API/agent integration suites, frontend Vitest, npm run typecheck and npm run build. Shared-schema/safety changes require full pytest regression and all frontend checks. Add actual browser workflow coverage: export/import or source-text assertions alone do not establish usability.')
p('Proposed performance targets for the demo: deterministic local query p95 under 3 seconds, alert visible within 5 seconds of a replay update, and bounded provider/LLM timeout with explicit fallback. Measure on the agreed demo machine; these are acceptance targets, not current performance claims.')

page('15. Release definition and remaining risks')
h('Fully functional prototype: definition of done')
bullets([
    'All eight supplied PS example query families run through real computations on declared data, with inspectable plans, evidence, maps/charts and uncertainty.',
    'All five stakeholder dashboards perform distinct useful workflows; role selection is backed by appropriate server-side scope for persistent operations.',
    'PFZ, risk, geofence and route tools consume canonical inputs. Changing a forecast, time, vessel or boundary changes the output when the rules require it.',
    'Synthetic seeding is reproducible and isolated. Demo time and real time are explicit; no operational fallback silently invents telemetry or availability.',
    'An adverse event or approaching boundary generates a proactive alert; deduplication, acknowledgement and incident/run linkage are demonstrable.',
    'A fresh checkout can start the complete demo, seed it and pass the release checks. Restart and provider-outage rehearsals preserve useful, honest behavior.'
])
table(['Risk', 'Mitigation / accountable role'], [
('Unverified government / EO access', 'M2 verifies source terms and sample access; retain labeled synthetic/historical mode. Do not let access delays block deterministic prototype work.'),
('Fixed demo time vs wall clock', 'M2/M4 publish one clock policy; M3 propagates it; M1 displays it. Never refresh timestamps merely to make fixtures appear current.'),
('False route safety / spatial coverage', 'M4 validates coastal geometry and time-aware route constraints; restrict to a tested pilot area.'),
('UI creates false operational confidence', 'M1 removes silent mocks and misleading live/health claims; M2 publishes actual availability and coverage.'),
('Shared mutable graph / scenario state', 'M3 isolates registries and memory; M2 tests overlapping API requests and restart behavior.'),
('Four-role schedule pressure', 'M2 gates integration; freeze decorative work and defer external messaging/national scope before sacrificing correctness.'),
('Evidence gaps mistaken for causal analysis', 'M4 defines scientific limits; M3 states hypotheses and required evidence; M1 shows coverage/QC explicitly.')
])
h('After the prototype')
p('A field pilot requires additional operational validation: authoritative boundary/source review, privacy/consent and access policies, load testing, monitoring/on-call response, backup/restore, user comprehension trials and marine-domain review. National routing, certified navigation, rescue dispatch, hardware AIS/NavIC and VHF/satellite integration are separate projects. Keep the existing prototype disclaimer visible.')
h('Immediate execution order')
p('First approve G0 semantics and file ownership. Then restore reproducible dependencies and establish the real test baseline. Close mock wiring, location/time substitution and outage-state defects before building the two new dashboards on top of those services.')

page('16. Repository evidence index')
p('Paths below are relative to the audited SAMUDRA repository at commit bf8e15a. Symbols are the primary locators because later work will shift line numbers. Findings marked present are source-inspection findings unless Section 1 explicitly reports execution.')
table(['Anchor', 'Files / symbols and evidence'], [
('E1 · Stakeholder surfaces', 'frontend/src/App.tsx: PortalMode and page switch; pages/FisherPage.tsx, AuthorityPage.tsx, ResearcherPage.tsx; domain/synthetic/generator.py: generate_stakeholders.'),
('E2 · Dispatch', 'backend/app/services/agent_run_service.py: run_agent invokes tool_mode="provider"; its explicit-instance LIVE/HYBRID guard differs from legacy service comments.'),
('E3 · Registration', 'backend/app/agents/graph.py: run_orca_graph; connectors/registration.py: register_dev2_provider_tools; main.py: create_app. Provider registration followed by register_m2_contract_mocks(override=False).'),
('E4 · Domain mocks / routes', 'agents/integrations/mocks.py: MockPFZRankingEngine, MockRouteExposureEngine, MockGeospatialHazardEngine, register_m2_contract_mocks; api/v1/routes.py: get_demo_route_alternatives directly uses MockRouteExposureEngine.'),
('E5 · Snapshot selection', 'connectors/snapshot.py: get_marine_conditions, _load_osf_fixture, get_svas_advisories. Fixed reference-row preference; Ratnagiri fallback; request-harbor reassignment; SAFE SVAS with newly generated validity.'),
('E6 · Safety / hazards', 'domain/risk_engine.py: DeterministicRiskEngine.evaluate, early is_data_degraded return; connectors/imd_hazard.py: _make_normal_payload and defaulted live timestamps; docs/SAFETY.md hard-stops.'),
('E7 · What-if / context', 'frontend/src/hooks/useChat.ts: simulateWhatIf; components/mission/WhatIfSimulator.tsx: handleRunSimulation and handleApply; agents/graph.py: specialist_tools_node arguments.'),
('E8 · UI fallbacks', 'frontend/src/api/researcher-client.ts: fetchOrMock, MOCK_MARINE_OBS, MOCK_HEALTH, fetchPFZCandidates, fetchHazards; pages/FisherPage.tsx baseline fetches.'),
('E9 · Analytics / agents', 'agents/supervisor.py: STANDARD_TOOL_PLANS; integrations/contracts.py capability catalog; agents/graph.py: supervisor_node; researcher components EOGridSpatialMap, EOTemporalAnalysisChart, TideTimeSeriesChart, QueryWorkbench.'),
('E10 · Alerts', 'domain/situation.py: get_canonical_active_hazards_for_sector and containment/operational-alert functions; api/v1/routes.py: sector alert endpoints; DECISIONS D018/D019 explicitly exclude scheduler and lifecycle.'),
('E11 · Seeds', 'domain/synthetic/generator.py, validator.py, seeder.py: entity_model_map and namespace rewrite; scripts/seed_demo.py; data/fixtures/synthetic; docs/SYNTHETIC_DATA.md.'),
('E12 · Platform', 'db/models.py, repositories.py, offline_store.py; backend/alembic/versions/d2f13393badb_initial_schema_with_postgis.py; main.py create_all; docker-compose.yml; backend/requirements.txt.'),
('E13 · Real domain foundation', 'domain/pfz.py: DeterministicPFZRankingEngine; domain/geo_restrictions.py: DeterministicGeospatialEngine; risk_engine.py; tests/domain, tests/api, tests/contract, tests/agent_eval.'),
('E14 · Governance / PS', 'Supplied PS 26176 pasted-text.txt; AGENTS.md, SKILLS.md, README.md; docs/OWNERSHIP.md, API_CONTRACTS.md, CANONICAL_DATA_CONTRACTS.md, SAFETY.md, ORCA_MASTER_CONTEXT.md, IMPLEMENTATION_PLAN.md, PROGRESS.md, DECISIONS.md.')
])
p('This fresh plan belongs in the requested report. After the proposed decisions are adopted, update the existing canonical implementation plan and progress board; do not create a competing roadmap.')

for row in doc.tables[-1].rows:
    for cell in row.cells:
        for para in cell.paragraphs:
            para.paragraph_format.space_before = Pt(1)
            para.paragraph_format.space_after = Pt(2)
            for run in para.runs:
                run.font.size = Pt(8.5)

path = OUT / 'ORCA_PS26176_Implementation_Audit_and_Plan.docx'
doc.save(path)
print(path)
print('Paragraphs:', len(doc.paragraphs), 'Tables:', len(doc.tables))
